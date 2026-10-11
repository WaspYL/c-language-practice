# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def add_h(text, level=1):
    doc.add_heading(text, level=level)

def add_p(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(11)

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text).font.size = Pt(11)

def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('CH32V307 项目接线与恢复指南')
r.bold = True
r.font.size = Pt(20)

doc.add_paragraph()

# 第1步：关机拔线
add_h('一、关机拔线顺序（断电操作）', 1)
add_p('每次用完按这个顺序，不要先拔信号：', bold=True)
add_bullet('1. 关掉PowerShell/串口助手窗口')
add_bullet('2. 拔掉板子上的Type-C电源线（先断电）')
add_bullet('3. 拔掉YD-LINK的USB头')
add_bullet('4. 网线可以拔也可以留')
add_bullet('5. 杜邦线建议保留插在模块上不动，下次直接插板子')

add_p('注意：所有接线改动都要先断电！', bold=True)

# 第2步：下次开机接线
add_h('二、下次开机接线（全断电接）', 1)
add_h('2.1 YD-LINK调试器', 2)
add_table(['线色', '接板子排针', '说明'], [
    ['红', '3.3V', '供电'],
    ['黄', 'SWDIO', '数据'],
    ['绿', 'SWCLK', '时钟'],
    ['黑', 'GND', '接地'],
])

add_h('2.2 ADXL345加速度传感器', 2)
add_table(['ADXL345引脚', '接板子', '说明'], [
    ['VCC', '3.3V', '电源'],
    ['GND', 'GND（面包板蓝轨）', '接地'],
    ['SCL', 'PB3', 'I2C时钟线'],
    ['SDA', 'PB5', 'I2C数据线'],
    ['CS', '3.3V', '接高=I2C模式'],
    ['SDO', 'GND', '接地=设备地址0x53'],
])

add_h('2.3 蜂鸣器', 2)
add_table(['蜂鸣器引脚', '接板子'], [
    ['VCC(+)', '3.3V'],
    ['GND(-)', 'GND'],
    ['I/O(信号)', 'PB1'],
])

add_h('2.4 其他', 2)
add_bullet('网线：板子网口 ↔ 电脑网口')
add_bullet('Type-C电源线：最后插，给板子供电')

# 第3步：电脑端操作
add_h('三、电脑端每次要做的事', 1)
add_bullet('1. 确认网口静态IP已设为192.168.1.10（之前设过，一般不用改）')
add_bullet('2. 开PowerShell，运行：python C:\\Users\\WaspYL\\Desktop\\c-language-practice\\test_server.py')
add_bullet('3. 按板子上的RST复位键')
add_bullet('4. 看到打印X,Y,Z,shake,level就成功了')

# 第4步：重新下载代码
add_h('四、重新下载代码', 1)
add_bullet('1. MRS打开工程：C:\\Users\\WaspYL\\Desktop\\c-language-practice\\MRS代码\\CH32V307_TCP')
add_bullet('2. 点编译（锤子图标）')
add_bullet('3. 点下载（向下箭头）')
add_bullet('4. 如果报错"找不到hex"：打开obj文件夹，把CH32V307_TCP.hex复制一份改名为TCPClient.hex')

# 第5步：常见问题
add_h('五、常见问题排查', 1)
add_table(['问题', '解决方法'], [
    ['下载失败', '按住BOOT0键不松→点下载→开始烧录后松开'],
    ['Python不是内部命令', '关掉PowerShell重新开一个新窗口'],
    ['板子连不上电脑', '检查网线、电脑IP是不是192.168.1.10'],
    ['数据不动', '按板子RST复位键'],
    ['LED不亮/屏不亮', '检查接线有没有松动，断电重插'],
])

doc.save(r'C:\Users\WaspYL\Desktop\c-language-practice\CH32V307项目接线与恢复指南.docx')
print('done')
